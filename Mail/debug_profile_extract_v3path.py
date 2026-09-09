# -*- coding: utf-8 -*-
import requests
import openpyxl

path = "ホーチミン企業リスト_全1403社.xlsx"
wb = openpyxl.load_workbook(path, data_only=True)
ws = wb.active

r = 2
name = ws.cell(r, 2).value
profile_url = ws.cell(r, 6).value

search_url = "https://www.fact-link.com.vn/search_result.php?id=003&type=prov&start=0&lang=jp"

headers = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36",
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
    "Accept-Language": "ja,en-US;q=0.7,en;q=0.3",
}

out_lines = []
out_lines.append(f"company={name}")
out_lines.append(f"profile_url={profile_url}")

s = requests.Session()
try:
    s.cookies.set("PHPSESSID", "auto", domain="www.fact-link.com.vn")
except Exception:
    pass

resp_search = s.get(search_url, headers=headers, timeout=30)
resp_search.encoding = "utf-8"
out_lines.append(f"search_status={resp_search.status_code}")
out_lines.append(f"search_len={len(resp_search.text)}")

ref_headers = dict(headers)
ref_headers["Referer"] = search_url

resp = s.get(profile_url, headers=ref_headers, timeout=30)
resp.encoding = "utf-8"
text = resp.text.strip()

out_lines.append(f"profile_status={resp.status_code}")
out_lines.append(f"profile_len={len(text)}")
out_lines.append(f"contains_refresh_meta={('refresh' in text.lower() and 'meta' in text.lower())}")
out_lines.append(f"contains_at_sign={('@' in text)}")

out_path = "debug_profile_extract_v3path.txt"
with open(out_path, "w", encoding="utf-8") as f:
    f.write("\n".join(out_lines))

print("wrote:", out_path)