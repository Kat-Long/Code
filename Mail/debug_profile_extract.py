# -*- coding: utf-8 -*-
import requests
from bs4 import BeautifulSoup
import re

path = "ホーチミン企業リスト_全1403社.xlsx"

import openpyxl
wb = openpyxl.load_workbook(path, data_only=True)
ws = wb.active

r = 2
name = ws.cell(r, 2).value
profile_url = ws.cell(r, 6).value
print("company:", name)
print("profile_url:", profile_url)

headers = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36",
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
    "Accept-Language": "ja,en-US;q=0.7,en;q=0.3",
    "Referer": "https://www.fact-link.com.vn/search_result.php?id=003&type=prov&start=0&lang=jp",
}

s = requests.Session()
resp = s.get(profile_url, headers=headers, timeout=30)
print("status_code:", resp.status_code)
resp.encoding = "utf-8"

text = resp.text
print("len(content):", len(text))
print("contains_refresh_meta:", ("refresh" in text.lower() and "meta" in text.lower()))
print("contains_at_sign:", ("@" in text))

soup = BeautifulSoup(resp.content, "html.parser")
page_text = soup.get_text("\n", strip=True)

# メール抽出（見えるテキストから）
email_pattern = r"[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}"
emails = re.findall(email_pattern, page_text)
print("emails_found:", len(emails))
if emails:
    print("first_email:", emails[0])

# 住所っぽい行（見えるテキストから短く）
hits = []
for ln in page_text.splitlines():
    ln = ln.strip()
    if not ln:
        continue
    if ("住所" in ln) or ("メール" in ln) or ("Ho Chi Minh" in ln) or ("ホーチミン" in ln):
        hits.append(ln)
        if len(hits) >= 5:
            break
print("address_like_lines_found:", len(hits))
for h in hits:
    print("line:", h)