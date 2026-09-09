# -*- coding: utf-8 -*-
import openpyxl

path = "ホーチミン企業リスト_全1403社.xlsx"
wb = openpyxl.load_workbook(path, data_only=True)
ws = wb.active

max_row = ws.max_row  # header included
total = max_row - 1

url_count = 0
email_count = 0
addr_count = 0

for r in range(2, max_row + 1):
    if ws.cell(r, 6).value:
        url_count += 1
    if ws.cell(r, 5).value:
        email_count += 1
    if ws.cell(r, 4).value:
        addr_count += 1

print(ws.title)
print(f"rows={total}")
print(f"url={url_count}/{total}")
print(f"addr={addr_count}/{total}")
print(f"email={email_count}/{total}")

# 先頭10件のサンプル（住所/メールの有無だけ）
shown = 0
for r in range(2, max_row + 1):
    if shown >= 10:
        break
    name = ws.cell(r, 2).value
    addr = ws.cell(r, 4).value
    email = ws.cell(r, 5).value
    print(f"{name} | addr={'Y' if addr else 'N'} | email={'Y' if email else 'N'}")
    shown += 1