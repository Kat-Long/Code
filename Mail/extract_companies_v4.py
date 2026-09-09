# -*- coding: utf-8 -*-
"""
FactLink Vietnam ホーチミン企業情報抽出スクリプト (v4)
短いタイムアウトで安定動作、新しいファイル名で保存
"""
import requests
from bs4 import BeautifulSoup
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
import time
import os

class FastScraper:
    def __init__(self):
        self.search_url_template = "https://www.fact-link.com.vn/search_result.php?id=003&type=prov&start={start}&lang=jp"
        self.headers = {
            'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36',
        }
        self.session = requests.Session()
        self.companies = []
    
    def fetch_page(self, start):
        url = self.search_url_template.format(start=start)
        try:
            r = self.session.get(url, headers=self.headers, timeout=20)
            r.encoding = 'utf-8'
            soup = BeautifulSoup(r.content, 'html.parser')
            
            results = []
            for span in soup.select('span.default_head'):
                a = span.find('a')
                if a and a.get('href'):
                    name = a.get_text(strip=True)
                    href = a['href']
                    full_url = 'https://www.fact-link.com.vn/' + href if href.startswith('mem') else href
                    
                    td = a.find_parent('td')
                    full_text = td.get_text(strip=True) if td else ''
                    biz = full_text.replace(name, '', 1).strip()
                    
                    results.append({
                        "name": name, "business": biz,
                        "address": "", "email": "",
                        "website": full_url
                    })
            return results
        except Exception as e:
            print(f"  [FAIL] start={start}: {str(e)[:50]}")
            return None  # None = timeout/error, empty list = no data
    
    def scrape_all(self, max_pages=52):
        print("ホーチミン企業リスト抽出 開始")
        print(f"全{max_pages}ページ対象\n")
        
        all_companies = []
        empty_count = 0
        
        for page in range(max_pages):
            start = page * 30
            print(f"  start={start:4d}... ", end="", flush=True)
            
            data = self.fetch_page(start)
            if data is None:
                print("TIMEOUT - リトライ")
                time.sleep(3)
                data = self.fetch_page(start)
            
            if data is None:
                print("SKIP")
                empty_count += 1
            elif len(data) == 0:
                print("0件（最終ページ）")
                empty_count += 1
                if empty_count >= 3:
                    print("\n3ページ連続0件のため終了")
                    break
            else:
                empty_count = 0
                all_companies.extend(data)
                print(f"{len(data):2d}件 (累計{len(all_companies):4d})")
            
            time.sleep(1)
        
        print(f"\n合計: {len(all_companies)} 件")
        self.companies = all_companies
        return all_companies
    
    def save_to_excel(self):
        path = u"c:\\Users\\long\\OneDrive\\Code\\Mail\\ホーチミン企業リスト_全1403社.xlsx"
        
        if os.path.exists(path):
            try:
                os.remove(path)
            except:
                path = u"c:\\Users\\long\\OneDrive\\Code\\Mail\\ホーチミン企業リスト_全社.xlsx"
        
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "企業リスト"
        
        headers = ["No.", "社名", "事業内容", "住所", "メール", "FactLinkウェブサイト"]
        ws.append(headers)
        
        hf = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        hfont = Font(bold=True, color="FFFFFF", size=11)
        for cell in ws[1]:
            cell.fill = hf
            cell.font = hfont
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        
        for i, c in enumerate(self.companies, 1):
            ws.append([i, c["name"], c["business"], "", "", c["website"]])
        
        ws.column_dimensions['A'].width = 5
        ws.column_dimensions['B'].width = 30
        ws.column_dimensions['C'].width = 45
        ws.column_dimensions['F'].width = 65
        
        for row in ws.iter_rows(min_row=2, max_row=len(self.companies)+1):
            for cell in row:
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                cell.font = Font(size=10)
        
        wb.save(path)
        print(f"\n[完了] {path}")
        print(f"[完了] {len(self.companies)} 社")

if __name__ == "__main__":
    s = FastScraper()
    s.scrape_all(max_pages=52)
    s.save_to_excel()