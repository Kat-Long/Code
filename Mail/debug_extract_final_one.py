# -*- coding: utf-8 -*-
import openpyxl
from extract_final import Scraper

path = "ホーチミン企業リスト_全1403社.xlsx"
wb = openpyxl.load_workbook(path, data_only=True)
ws = wb.active

r = 2
profile_url = ws.cell(r, 6).value
name = ws.cell(r, 2).value

s = Scraper()

out_path = "debug_extract_final_one.txt"

try:
    # session確立（v3同等に start=0 検索をReferer付きで1回）
    try:
        s.session.cookies.set("PHPSESSID", "auto", domain="www.fact-link.com.vn")
    except Exception:
        pass

    try:
        s.session.get(s.search_referer, headers=s.headers, timeout=20)
    except Exception:
        pass

    res = s.try_fetch_profile(profile_url)

    addr = res.get("address", "") or ""
    email = res.get("email", "") or ""

    with open(out_path, "w", encoding="utf-8") as f:
        f.write(f"company={name}\n")
        f.write(f"profile_url={profile_url}\n")
        f.write(f"address_len={len(addr)}\n")
        f.write(f"email_len={len(email)}\n")
        f.write(f"address_prefix={addr[:120]}\n")
        f.write(f"email_prefix={email[:120]}\n")

    print("wrote:", out_path)

except Exception as e:
    with open(out_path, "w", encoding="utf-8") as f:
        f.write(f"company={name}\n")
        f.write(f"profile_url={profile_url}\n")
        f.write(f"ERR={type(e).__name__}: {e}\n")
    print("wrote error:", out_path)