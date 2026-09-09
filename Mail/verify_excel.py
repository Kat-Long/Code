# -*- coding: utf-8 -*-
import openpyxl

path = "ホーチミン企業リスト_全1403社.xlsx"

wb = openpyxl.load_workbook(path, data_only=True)
ws = wb.active

print(f"Sheet: {ws.title}")
print(f"Total rows (incl header): {ws.max_row}")
print(f"Total companies: {ws.max_row - 1}")
print(f"Headers: {[ws.cell(1, c).value for c in range(1, 7)]}")
print(f"First: {ws.cell(2, 2).value}")
print(f"Last: {ws.cell(ws.max_row, 2).value}")

url_count = 0
email_count = 0
for r in range(2, ws.max_row + 1):
    if ws.cell(r, 6).value:
        url_count += 1
    if ws.cell(r, 5).value:
        email_count += 1

print(f"Companies with URLs: {url_count}")
print(f"Companies with Emails: {email_count}")