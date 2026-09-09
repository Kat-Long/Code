# -*- coding: utf-8 -*-
import sys
from pathlib import Path

import openpyxl


def main():
    excel_path = "ホーチミン企業リスト_全1403社.xlsx"
    if len(sys.argv) >= 2:
        excel_path = sys.argv[1]

    excel_path = Path(excel_path)
    if not excel_path.exists():
        print(f"not_found: {excel_path}")
        return 2

    wb = openpyxl.load_workbook(str(excel_path), data_only=True)
    ws = wb.active

    # verify column indices used elsewhere:
    # url: col 6, addr: col 4, email: col 5
    url_col = 6
    addr_col = 4
    email_col = 5

    max_row = ws.max_row
    if max_row < 2:
        print("no_data")
        return 0

    missing_addr = 0
    missing_email = 0
    missing_url = 0
    total = max_row - 1

    # Collect a few sample rows for quick diagnostics
    samples = []
    for r in range(2, max_row + 1):
        url_v = ws.cell(r, url_col).value
        addr_v = ws.cell(r, addr_col).value
        email_v = ws.cell(r, email_col).value

        if not url_v:
            missing_url += 1
        if not addr_v:
            missing_addr += 1
        if not email_v:
            missing_email += 1

        if len(samples) < 10 and (not url_v or not addr_v or not email_v):
            # attempt to include company name if present (best effort)
            name_v = ws.cell(r, 1).value  # may or may not exist
            samples.append((r, name_v, addr_v, email_v, url_v))

    print(f"rows={total}")
    print(f"missing_url(col{url_col})={missing_url}/{total}")
    print(f"missing_addr(col{addr_col})={missing_addr}/{total}")
    print(f"missing_email(col{email_col})={missing_email}/{total}")

    if samples:
        print("samples(row,company,addr,email,url):")
        for (r, name_v, addr_v, email_v, url_v) in samples:
            # keep output single-line and compact
            name_s = "" if name_v is None else str(name_v)[:80]
            addr_s = "" if addr_v is None else str(addr_v)[:80]
            email_s = "" if email_v is None else str(email_v)[:80]
            url_s = "" if url_v is None else str(url_v)[:120]
            print(f"{r}\t{name_s}\t{addr_s}\t{email_s}\t{url_s}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())